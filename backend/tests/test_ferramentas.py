from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.ferramentas import criar_apresentacao, criar_documento, criar_planilha


def test_criar_documento(tmp_path):
    caminho = criar_documento(tmp_path, "Relatório", [{"titulo": "Resumo", "paragrafos": ["Texto"]}])
    textos = [p.text for p in Document(caminho).paragraphs]
    assert textos == ["Relatório", "Resumo", "Texto"]


def test_criar_planilha(tmp_path):
    caminho = criar_planilha(tmp_path, "Orçamento", ["Item", "Valor"], [["Anúncios", 1200]])
    ws = load_workbook(caminho).active
    assert [c.value for c in ws[1]] == ["Item", "Valor"]
    assert [c.value for c in ws[2]] == ["Anúncios", 1200]


def test_criar_apresentacao(tmp_path):
    caminho = criar_apresentacao(tmp_path, "Plano", [{"titulo": "Metas", "topicos": ["A", "B"]}])
    prs = Presentation(caminho)
    assert len(prs.slides) == 2
    assert prs.slides[1].placeholders[1].text_frame.text == "A\nB"
