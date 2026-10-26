from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.exportar import documento_para_docx
from app.main import app

DOCUMENTO = {
    "type": "doc",
    "content": [
        {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Relatório"}]},
        {
            "type": "paragraph",
            "content": [
                {"type": "text", "text": "Vendas "},
                {"type": "text", "text": "subiram", "marks": [{"type": "bold"}]},
            ],
        },
        {"type": "paragraph"},
        {
            "type": "bulletList",
            "content": [
                {"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Item"}]}]}
            ],
        },
        {
            "type": "table",
            "content": [
                {"type": "tableRow", "content": [
                    {"type": "tableHeader", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Mês"}]}]},
                    {"type": "tableHeader", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Total"}]}]},
                ]},
                {"type": "tableRow", "content": [
                    {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Jan"}]}]},
                    {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "10"}]}]},
                ]},
            ],
        },
    ],
}


def test_documento_para_docx():
    doc = Document(BytesIO(documento_para_docx(DOCUMENTO)))
    paragrafos = [(p.style.name, p.text) for p in doc.paragraphs]
    assert paragrafos == [("Heading 1", "Relatório"), ("Normal", "Vendas subiram"), ("Normal", ""), ("List Bullet", "Item")]
    assert doc.paragraphs[1].runs[1].bold is True
    tabela = doc.tables[0]
    assert [c.text for c in tabela.rows[1].cells] == ["Jan", "10"]
    assert tabela.cell(0, 0).paragraphs[0].runs[0].bold is True


def test_endpoint_exportar_docx():
    resposta = TestClient(app).post("/api/exportar/docx", json={"documento": DOCUMENTO, "nome": "Relatório 2026"})
    assert resposta.status_code == 200
    assert "Relat%C3%B3rio%202026.docx" in resposta.headers["content-disposition"]
    assert Document(BytesIO(resposta.content)).paragraphs[0].text == "Relatório"


def test_endpoint_sugestoes_exige_texto():
    resposta = TestClient(app).post("/api/sugestoes", json={"paragrafo": "  "})
    assert resposta.status_code == 400
