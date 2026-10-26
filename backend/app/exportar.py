"""Converte o documento do editor (JSON do TipTap/ProseMirror) para .docx."""

from io import BytesIO

from docx import Document

ESTILOS_LISTA = {"bulletList": "List Bullet", "orderedList": "List Number"}


def documento_para_docx(documento: dict) -> bytes:
    doc = Document()
    for no in documento.get("content", []):
        _adicionar_bloco(doc, no)
    saida = BytesIO()
    doc.save(saida)
    return saida.getvalue()


def _adicionar_bloco(doc, no: dict) -> None:
    tipo = no.get("type")
    if tipo == "heading":
        nivel = no.get("attrs", {}).get("level", 1)
        _adicionar_texto(doc.add_heading(level=nivel), no)
    elif tipo == "paragraph":
        _adicionar_texto(doc.add_paragraph(), no)
    elif tipo in ESTILOS_LISTA:
        for item in no.get("content", []):
            for filho in item.get("content", []):
                _adicionar_texto(doc.add_paragraph(style=ESTILOS_LISTA[tipo]), filho)
    elif tipo == "blockquote":
        for filho in no.get("content", []):
            _adicionar_texto(doc.add_paragraph(style="Quote"), filho)
    elif tipo == "table":
        _adicionar_tabela(doc, no)


def _adicionar_texto(paragrafo, no: dict) -> None:
    for trecho in no.get("content", []):
        if trecho.get("type") == "hardBreak":
            paragrafo.add_run().add_break()
            continue
        run = paragrafo.add_run(trecho.get("text", ""))
        marcas = {m.get("type") for m in trecho.get("marks", [])}
        run.bold = "bold" in marcas or None
        run.italic = "italic" in marcas or None


def _adicionar_tabela(doc, no: dict) -> None:
    linhas = no.get("content", [])
    colunas = max((len(linha.get("content", [])) for linha in linhas), default=0)
    if not linhas or not colunas:
        return
    tabela = doc.add_table(rows=len(linhas), cols=colunas)
    tabela.style = "Table Grid"
    for i, linha in enumerate(linhas):
        for j, celula in enumerate(linha.get("content", [])):
            destino = tabela.cell(i, j)
            destino.text = ""
            blocos = celula.get("content", [])
            for k, bloco in enumerate(blocos):
                paragrafo = destino.paragraphs[0] if k == 0 else destino.add_paragraph()
                _adicionar_texto(paragrafo, bloco)
            if celula.get("type") == "tableHeader":
                for run in destino.paragraphs[0].runs:
                    run.bold = True
