from pathlib import Path

from docx import Document

from .util import caminho_saida

SCHEMA = {
    "type": "function",
    "function": {
        "name": "criar_documento",
        "description": "Cria um documento de texto (.docx) com título e seções.",
        "parameters": {
            "type": "object",
            "properties": {
                "titulo": {"type": "string"},
                "secoes": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "titulo": {"type": "string"},
                            "paragrafos": {"type": "array", "items": {"type": "string"}},
                        },
                        "required": ["titulo", "paragrafos"],
                    },
                },
            },
            "required": ["titulo", "secoes"],
        },
    },
}


def criar_documento(pasta: Path, titulo: str, secoes: list[dict]) -> Path:
    doc = Document()
    doc.add_heading(titulo, level=0)
    for secao in secoes:
        doc.add_heading(secao.get("titulo", ""), level=1)
        for paragrafo in secao.get("paragrafos", []):
            doc.add_paragraph(paragrafo)
    caminho = caminho_saida(pasta, titulo, "docx")
    doc.save(caminho)
    return caminho
