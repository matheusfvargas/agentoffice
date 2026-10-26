from pathlib import Path

from pptx import Presentation

from .util import caminho_saida

SCHEMA = {
    "type": "function",
    "function": {
        "name": "criar_apresentacao",
        "description": "Cria uma apresentação (.pptx) com um slide de capa e slides de tópicos.",
        "parameters": {
            "type": "object",
            "properties": {
                "titulo": {"type": "string"},
                "slides": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "titulo": {"type": "string"},
                            "topicos": {"type": "array", "items": {"type": "string"}},
                        },
                        "required": ["titulo", "topicos"],
                    },
                },
            },
            "required": ["titulo", "slides"],
        },
    },
}


def criar_apresentacao(pasta: Path, titulo: str, slides: list[dict]) -> Path:
    prs = Presentation()
    capa = prs.slides.add_slide(prs.slide_layouts[0])
    capa.shapes.title.text = titulo
    for dados in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = dados.get("titulo", "")
        corpo = slide.placeholders[1].text_frame
        topicos = dados.get("topicos", [])
        if topicos:
            corpo.text = topicos[0]
            for topico in topicos[1:]:
                corpo.add_paragraph().text = topico
    caminho = caminho_saida(pasta, titulo, "pptx")
    prs.save(caminho)
    return caminho
